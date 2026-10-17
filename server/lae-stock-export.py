import sys,json,io,datetime,html
report=json.load(sys.stdin);models=report['filters'].get('view')=='models'
batch_columns=[('productGroup','Item group'),('masterCode','Master code'),('brandCode','Brand code'),('itemName','Item name / model'),('batch','Barcode / batch'),('warehouse','Warehouse'),('unit','Unit'),('quantity','Closing quantity'),('assignedRo','Assigned RO'),('originalRo','Recorded RO'),('priceRo','Price-reference RO'),('invoice','Invoice'),('supplier','Supplier'),('usdUnit','USD per unit'),('originalRate','Original INR per USD'),('originalInrUnit','Original landing INR per unit'),('interestInrUnit','Interest INR per unit'),('rentInrUnit','Rent INR per unit'),('dutyInrUnit','Duty addition INR per unit'),('revisedRate','Revised INR per USD'),('revisedInrUnit','Revised INR per unit'),('changePercent','Change percent'),('status','Status'),('rentExcluded','Approved rent exclusion'),('quantityOnly','Quantity-only stock'),('inwardDate','Warehouse inward date'),('importDate','Import-age date proxy'),('sourceSheet','Source sheet'),('sourceRow','Source row'),('calcId','Calculation ID'),('sourceEvidence','Cost evidence'),('matchBasis','RO matching basis'),('validationConcerns','Review notes'),('motorOriginalInrUnit','Separate historical motor INR'),('saleSetOriginalInrUnit','Historical sale-set original INR')]
model_columns=[('productGroup','Item group'),('masterCode','Master code'),('itemName','Item name / model'),('brand','Brands'),('warehouse','Warehouses'),('unit','Unit'),('netQuantity','Net closing quantity'),('positiveQuantity','Positive quantity'),('negativeQuantity','Negative adjustments'),('averageUsd','Weighted USD'),('originalRate','Original INR per USD'),('originalAverage','Original average INR'),('originalPricedQuantity','Original priced quantity'),('originalPendingQuantity','Original pending quantity'),('revisedRate','Revised INR per USD'),('revisedAverage','Revised average INR'),('revisedPricedQuantity','Revised priced quantity'),('revisedPendingQuantity','Revised pending quantity'),('comparisonQuantity','Compared quantity'),('comparisonOriginalAverage','Original average on compared stock'),('comparisonRevisedAverage','Revised average on compared stock'),('changePercent','Change percent on same stock'),('quantityOnlyQuantity','Quantity-only stock'),('rentExcludedQuantity','Approved rent-excluded quantity')]
model_columns += [('importBatchCount','Import batches (ROs)'),('importRos','Purchase RO numbers'),('roPendingQuantity','RO pending quantity')]
batch_columns += [('purchaseRo','Purchase RO'),('roAssignmentStatus','RO assignment status')]
columns=model_columns if models else batch_columns;snapshot=report.get('snapshot') or {};source=snapshot.get('source') or {}
meta='Closing stock: '+snapshot.get('stockDate','Pending')+' | Costs: '+snapshot.get('costDate','Pending')+' | Before GST | Physical stock'
filter_text='; '.join(str(k)+': '+str(v) for k,v in report['filters'].items() if k not in ['page','pageSize','sort','direction'] and v!='')
def cell_value(value):
 if isinstance(value,list):value='; '.join(str(v) for v in value)
 if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):return "'"+value
 return value
if sys.argv[1]=='xlsx':
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.worksheet.table import Table,TableStyleInfo
 from openpyxl.utils import get_column_letter
 w=Workbook();s=w.active;s.title='Model costs' if models else 'Batch costs';s.append(['LAE closing stock costing']);s.append([meta]);s.append(['Filters: '+filter_text]);s.append(['Source: '+source.get('filename','')+' | SHA256: '+source.get('sha256','')]);s.append(['Averages exclude missing prices and non-positive quantity. Approved rent exclusions stay visible.']);s.append([v for _,v in columns])
 for row in report['rows']:s.append([cell_value(row.get(k)) for k,_ in columns])
 for c in s[6]:c.fill=PatternFill('solid',fgColor='204321');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 s.row_dimensions[6].height=44;s.freeze_panes='F7';s.sheet_view.showGridLines=False;s.auto_filter.ref='A6:'+get_column_letter(len(columns))+str(max(6,s.max_row))
 if report['rows']:
  table=Table(displayName='ClosingStockCosting',ref=s.auto_filter.ref);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);s.add_table(table)
 for i,(key,label) in enumerate(columns,1):s.column_dimensions[get_column_letter(i)].width=34 if key in ['itemName','warehouse','batch'] else 46 if key in ['sourceEvidence','matchBasis','validationConcerns'] else 20
 for row in s.iter_rows(min_row=7):
  for c in row:
   if isinstance(c.value,(int,float)) and not isinstance(c.value,bool):c.number_format='#,##0.00'
 s.print_title_rows='1:6';s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
 t=w.create_sheet('Selected totals');t.append(['Selected closing stock totals']);t.append([meta]);t.append(['Same filters as the detail sheet. Positive closing quantities weight prices.']);t.append([])
 totals=[('unit','Unit'),('netQuantity','Net stock'),('originalAverage','Original avg INR'),('originalPricedQuantity','Original priced qty'),('originalPendingQuantity','Original pending qty'),('revisedAverage','Revised avg INR'),('revisedPricedQuantity','Revised priced qty'),('revisedPendingQuantity','Revised pending qty'),('comparisonQuantity','Compared qty'),('changePercent','Change percent / same stock')]
 totals += [('importBatchCount','Import batches (ROs)'),('roPendingQuantity','RO pending quantity')]
 t.append([v for _,v in totals]);
 for r in report['summary']['byUnit']:t.append([cell_value(r.get(k)) for k,_ in totals])
 t.append([]);t.append(['Warehouse']+[v for _,v in totals]);
 for r in report.get('warehouses') or []:t.append([cell_value(r['warehouse'])]+[cell_value(r.get(k)) for k,_ in totals])
 t.freeze_panes='B6';t.sheet_view.showGridLines=False
 for col in range(1,len(totals)+2):t.column_dimensions[get_column_letter(col)].width=27
 for row in t:
  for c in row:
   if isinstance(c.value,(int,float)):c.number_format='#,##0.00'
 r=w.create_sheet('Import batches');r.append(['Selected import batches (distinct RO): '+str(report['summary'].get('importBatchCount',0))]);r.append([meta]);r.append(['Counts use positive-stock RO assignments; unknown or multiple assignments stay pending.']);r.append(['Purchase RO','Unit','Net closing stock','Original average INR','Revised average INR']);
 for row in report.get('importBatches') or []:r.append([cell_value(row.get('ro') or 'RO pending'),cell_value(row.get('unit')),row.get('netQuantity'),row.get('originalAverage'),row.get('revisedAverage')])
 r.freeze_panes='C5';r.auto_filter.ref='A4:E'+str(r.max_row);
 for col in range(1,6):r.column_dimensions[get_column_letter(col)].width=25
 for row in r.iter_rows(min_row=5,min_col=3,max_col=5):
  for c in row:
   if isinstance(c.value,(int,float)):c.number_format='#,##0.00'
 output=io.BytesIO();w.save(output);sys.stdout.buffer.write(output.getvalue())
elif sys.argv[1]=='pdf':
 from reportlab.lib.pagesizes import A3,landscape
 from reportlab.lib import colors
 from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
 from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle
 out=io.BytesIO();doc=SimpleDocTemplate(out,pagesize=landscape(A3),leftMargin=24,rightMargin=24,topMargin=24,bottomMargin=26);styles=getSampleStyleSheet();style=ParagraphStyle('stockcell',fontName='Helvetica',fontSize=8,leading=10);story=[Paragraph('LAE closing stock costing',styles['Title']),Paragraph(html.escape(meta),styles['Normal']),Paragraph(html.escape(filter_text),styles['Normal']),Spacer(1,8)]
 for r in report['summary']['byUnit']:
  n=lambda v:'Pending' if v is None else f'{v:,.2f}'
  story.append(Paragraph(html.escape(f"{r['unit']}: net stock {n(r['netQuantity'])}; original average INR {n(r['originalAverage'])} ({n(r['originalPricedQuantity'])} priced); revised average INR {n(r['revisedAverage'])} ({n(r['revisedPricedQuantity'])} priced, {n(r['revisedPendingQuantity'])} pending). Approved rent-excluded quantity: {n(r['rentExcludedQuantity'])}."),styles['Normal']))
 story.append(Paragraph('Import batches (distinct ROs): '+str(report['summary'].get('importBatchCount',0))+'. Unknown or multiple RO assignments are excluded.',styles['Normal']));
 for r in report['summary']['byUnit']:
  story.append(Paragraph(html.escape((r.get('unit') or 'Unspecified unit')+' | Purchase ROs: '+(', '.join(r.get('importRos') or []) or 'Pending')+' | RO pending quantity: '+str(r.get('roPendingQuantity',0))),styles['Normal']))
 story.append(Spacer(1,8));keys=[('masterCode','Master'),('brandCode','Brand code'),('batch','Barcode / batch'),('warehouse','Warehouse'),('purchaseRo','Purchase RO'),('quantity','Qty'),('usdUnit','USD'),('originalInrUnit','Original INR'),('interestInrUnit','Interest'),('rentInrUnit','Rent'),('dutyInrUnit','Duty'),('revisedInrUnit','Revised INR')]
 if models:keys=[('productGroup','Group'),('masterCode','Master'),('itemName','Model / item'),('warehouse','Warehouses'),('netQuantity','Qty'),('averageUsd','Avg USD'),('originalAverage','Original avg INR'),('revisedAverage','Revised avg INR'),('revisedPendingQuantity','Pending qty')]
 def p(v):return Paragraph(html.escape('Pending' if v is None or v=='' else f'{v:,.2f}' if isinstance(v,(int,float)) else str(v)).replace('\n','<br/>'),style)
 body=[[p(label) for _,label in keys]]+[[p('0.00*' if k=='rentInrUnit' and r.get('rentExcluded') else r.get(k)) for k,_ in keys] for r in report['rows']];widths=([68,80,170,150,55,55,65,85,75,65,65,95] if not models else [110,70,210,240,65,80,100,100,95]);scale=(landscape(A3)[0]-48)/sum(widths);table=LongTable(body,colWidths=[v*scale for v in widths],repeatRows=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dce8dc')),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f1f5f1')]),('BOTTOMPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5)]));story.append(table);story.append(Spacer(1,8));story.append(Paragraph('Source: '+html.escape(source.get('filename',''))+'. Missing costs are excluded from averages. * Rent excluded under the approved unavailable-rent policy. Costs are dated; stock is not a live inventory feed.',styles['Normal']))
 def footer(canvas,document):canvas.setFont('Helvetica',8);canvas.drawRightString(landscape(A3)[0]-24,13,'Page '+str(document.page))
 doc.build(story,onFirstPage=footer,onLaterPages=footer);sys.stdout.buffer.write(out.getvalue())
else:raise ValueError('Unsupported report format')
