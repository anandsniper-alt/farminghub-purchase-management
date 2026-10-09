"""Filtered, before-GST purchase register using the existing application export runtime."""
import datetime as dt
import io
import json
import sys
import textwrap
from urllib.parse import quote
from xml.sax.saxutils import escape

COLUMNS = [('ro','RO #'),('inwardDate','Inward date'),('vendorName','Vendor name'),('itemName','Item name'),('masterCode','Master code'),('brandCode','Brand code'),('brand','Brand'),('usdUnit','USD / unit'),('aiInrPerUsd','AI INR / USD'),('inrUnitBeforeGst','INR / unit before GST'),('quantity','Qty'),('unit','Unit'),('status','Costing status'),('invoice','Supplier invoice')]
FILTER_LABELS = {'q':'Search','year':'Inward year','month':'Inward month','supplier':'Supplier','product':'Product group','masterCode':'Master code','brand':'Brand','unit':'Unit','status':'Costing status'}
SUMMARY_COLUMNS = [('totalUnit','Unit'),('rowCount','Rows'),('totalQuantity','Purchase qty'),('pairedQuantity','Costed qty'),('pendingQuantity','Pending qty'),('avgUsd','Avg USD / unit'),('effectiveInrPerUsd','INR / USD'),('avgInr','Avg INR / unit'),('totalInr','Costed total INR')]
def filter_text(data):
    filters=data.get('filters',{})
    pairs=[f'{label}: {"Pending" if filters.get(key)=="__missing__" else filters[key]}' for key,label in FILTER_LABELS.items() if filters.get(key)]
    return ' | '.join(pairs) if pairs else 'All recorded purchase items'
def generated(data):
    value=dt.datetime.fromisoformat(data['generatedAt'].replace('Z','+00:00')).astimezone(dt.timezone(dt.timedelta(hours=5,minutes=30)))
    return value.strftime('%d %b %Y, %I:%M %p IST')
def excel(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook();ws=wb.active;ws.title='Landing Prices';ws.sheet_view.showGridLines=False
    wb.calculation=CalcProperties(fullCalcOnLoad=True)
    green='204321'; light='EDF4E8';n=len(COLUMNS)
    widths=[14,17,39,49,18,22,19,17,18,22,14,12,25,31]
    def literal(cell,value):
        cell.value=value
        if isinstance(value,str):cell.data_type='s'
        cell.alignment=Alignment(vertical='top',wrap_text=True)
    for row,value in [(1,'Farming Hub | LAE Landing Prices'),(2,'Before GST | '+generated(data)),(3,filter_text(data)),(4,f'{len(data["rows"]):,} selected purchase lines | Snapshot of the selected website filters and sorting')]:
        ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=n);literal(ws.cell(row,1),value)
    ws['A1'].font=Font(name='Calibri',size=18,bold=True,color=green);ws.row_dimensions[1].height=28;ws.row_dimensions[3].height=34
    header=6
    for col,(_,label) in enumerate(COLUMNS,1):
        cell=ws.cell(header,col,label);cell.fill=PatternFill('solid',fgColor=green);cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[header].height=34
    for index,row in enumerate(data['rows'],header+1):
        for col,(key,_) in enumerate(COLUMNS,1):
            value=row.get(key)
            if value is None or value=='':value='Pending'
            if key=='inwardDate' and value!='Pending':value=dt.date.fromisoformat(value)
            cell=ws.cell(index,col);literal(cell,value);cell.font=Font(name='Calibri',size=11,color='203723')
            if key=='inwardDate' and value!='Pending':cell.number_format='dd-mmm-yyyy'
            elif key=='aiInrPerUsd' and isinstance(value,(int,float)):cell.number_format='#,##0.0000'
            elif key in ('usdUnit','inrUnitBeforeGst') and isinstance(value,(int,float)):cell.number_format='#,##0.00'
            elif key=='quantity' and isinstance(value,(int,float)):cell.number_format='#,##0.##'
            if key=='ro':cell.hyperlink='https://purchase.dvjassociates.com/#/ro-costings/'+quote(str(value),safe='');cell.font=Font(color=green,underline='single')
        lines=max(sum(max(1,len(textwrap.wrap(part,width=max(1,int(widths[col-1])-2),break_long_words=True,break_on_hyphens=False))) for part in str(ws.cell(index,col).value or '').split('\n')) for col in range(1,n+1))
        ws.row_dimensions[index].height=min(409,max(32,15*lines+8))
    end=header+len(data['rows'])
    if data['rows']:
        table=Table(displayName='LandingPrices',ref=f'A{header}:N{end}');table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);ws.add_table(table)
    else:literal(ws.cell(header+1,1),'No matching purchase lines')
    footer=max(end,header+1)+3
    ws.merge_cells(start_row=footer,start_column=1,end_row=footer,end_column=n);ws.cell(footer,1,'WEIGHTED AVERAGES - ALL SELECTED ROWS').font=Font(bold=True,color=green,size=13)
    for col,(_,label) in enumerate(SUMMARY_COLUMNS,1):
        ws.cell(footer+1,col,label).font=Font(bold=True,color=green);ws.cell(footer+1,col).alignment=Alignment(wrap_text=True);ws.cell(footer+1,col).fill=PatternFill('solid',fgColor=light)
    for index,row in enumerate(data['summary'].get('byUnit',[]),footer+2):
        for col,(key,_) in enumerate(SUMMARY_COLUMNS,1):
            literal(ws.cell(index,col),row.get(key) if row.get(key) is not None else 'Pending')
            if col>=6 and isinstance(ws.cell(index,col).value,(int,float)):ws.cell(index,col).number_format='#,##0.0000' if key=='effectiveInrPerUsd' else '#,##0.00'
    note=footer+3+len(data['summary'].get('byUnit',[]))
    notes=['Averages use positive quantities with both USD and supported INR costs. Conversion = paired INR value / paired USD value. Pieces and sets remain separate.', 'Missing costs are excluded, not zero. Historical and provisional references retain their recorded status. Downloaded values are a snapshot; filter again on the website for recalculated averages.']
    for offset,text in enumerate(notes):ws.merge_cells(start_row=note+offset,start_column=1,end_row=note+offset,end_column=n);literal(ws.cell(note+offset,1),text);ws.row_dimensions[note+offset].height=30
    for col,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(col)].width=width
    ws.freeze_panes='D7';ws.print_title_rows='1:6';ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    stream=io.BytesIO();wb.save(stream);return stream.getvalue()
def pdf(data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3,landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,Table,TableStyle,KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from pathlib import Path
    # Existing OS fonts, not an added package or external network dependency.
    font='Helvetica'
    for path in ['/usr/share/fonts/truetype/dejavu/dejavusans.ttf','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',r'C:\Windows\Fonts\arial.ttf']:
        if Path(path).exists():pdfmetrics.registerFont(TTFont('LandingSans',path));font='LandingSans';break
    green=colors.HexColor('#204321');light=colors.HexColor('#EDF4E8');stream=io.BytesIO();size=landscape(A3)
    doc=SimpleDocTemplate(stream,pagesize=size,leftMargin=28,rightMargin=28,topMargin=28,bottomMargin=34,title='LAE Landing Prices',author='Farming Hub')
    normal=ParagraphStyle('normal',fontName=font,fontSize=8,leading=11,textColor=green,wordWrap='CJK');small=ParagraphStyle('small',parent=normal,fontSize=7,leading=9);right=ParagraphStyle('right',parent=normal,alignment=TA_RIGHT)
    header=ParagraphStyle('header',parent=normal,textColor=colors.white,fontSize=8,leading=11)
    p=lambda value,style=normal:Paragraph(escape(str(value if value not in (None,'') else 'Pending')).replace('\n','<br/>'),style)
    fmt=lambda value,d=2: 'Pending' if value is None else f'{value:,.{d}f}'
    flow=[Paragraph('Farming Hub | LAE Landing Prices',ParagraphStyle('title',parent=normal,fontSize=19,leading=24)),Spacer(1,5),p('Before GST | '+generated(data)),p(filter_text(data)),p(f'{len(data["rows"]):,} selected purchase lines'),Spacer(1,12)]
    labels=['RO #','Inward date','Vendor / invoice','Item name','Master code','Brand code','Brand','USD / unit','AI INR / USD','INR / unit before GST','Qty / unit']
    widths=[49,69,168,223,68,86,76,73,82,99,65]
    rows=[[p(label,header) for label in labels]]
    for row in data['rows']:
        date=dt.date.fromisoformat(row['inwardDate']).strftime('%d %b %Y') if row.get('inwardDate') else 'Pending'
        code=escape(row['ro']);url='https://purchase.dvjassociates.com/#/ro-costings/'+quote(row['ro'],safe='')
        rows.append([Paragraph(f'<link href="{escape(url)}" color="#204321"><u>{code}</u></link>',normal),p(date),p(row['vendorName']+'\n'+row.get('invoice','')),p(row['itemName']),p(row['masterCode']),p(row['brandCode']),p(row['brand']),p(fmt(row['usdUnit']),right),p(fmt(row['aiInrPerUsd'],4),right),p(fmt(row['inrUnitBeforeGst'])+'\n'+row['status'],right),p(fmt(row['quantity'],0 if row['quantity'] is not None and float(row['quantity']).is_integer() else 2)+' '+row.get('unit',''),right)])
    table=LongTable(rows,colWidths=widths,repeatRows=1,hAlign='LEFT');table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),green),('VALIGN',(0,0),(-1,-1),'TOP'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F6F8F3')]),('LINEBELOW',(0,0),(-1,0),.6,green),('BOTTOMPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7)]));flow.append(table)
    if not data['rows']:flow.append(p('No matching purchase lines.'))
    summary=[[p(label,header) for _,label in SUMMARY_COLUMNS]]
    for row in data['summary'].get('byUnit',[]):summary.append([p(row.get(key) if key=='totalUnit' else fmt(row.get(key),4 if key=='effectiveInrPerUsd' else 2)) for key,_ in SUMMARY_COLUMNS])
    footer=Table(summary,colWidths=[70,60,100,100,100,115,115,130,160]);footer.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),green),('BACKGROUND',(0,1),(-1,-1),light),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),9)]))
    flow.append(KeepTogether([Spacer(1,16),p('WEIGHTED AVERAGES - ALL SELECTED ROWS'),Spacer(1,6),footer,Spacer(1,7),p('Positive quantities with both USD and supported INR costs. Conversion = paired INR value / paired USD value. Missing costs excluded; pieces and sets separate.',small)]))
    def page(canvas,document):
        canvas.setFont(font,8);canvas.setFillColor(green);canvas.drawString(28,18,'Farming Hub | Before GST | Filtered purchase-price snapshot');canvas.drawRightString(size[0]-28,18,f'Page {document.page}')
    doc.build(flow,onFirstPage=page,onLaterPages=page);return stream.getvalue()
if __name__=='__main__':
    try:
        data=json.load(sys.stdin)
        if len(data['rows'])>25000:raise ValueError('Too many rows')
        sys.stdout.buffer.write(excel(data) if sys.argv[1]=='xlsx' else pdf(data))
    except Exception as error:
        print(type(error).__name__,file=sys.stderr);sys.exit(1)
