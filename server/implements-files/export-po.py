"""Portable review PO workbook with quantity and fabrication schedules."""
import sys,json,io,base64
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
from PIL import Image as PILImage
po=json.load(sys.stdin);wb=Workbook();ws=wb.active;ws.title='Purchase Order'
rows=[['FARMING HUB - REVIEW PURCHASE ORDER',po['id']],['Status','Approval and issuing pending'],['Date',po['date']],['Supplier',po['supplier']['name']],['Supplier address',po['supplier'].get('address','')],['Buyer',po['buyer'].get('buyer','')],['Buyer address',po['buyer'].get('address','')],['Delivery',po.get('delivery','')],['Planning month',po.get('planningMonth','')],[],['Item code','Component','MRP pcs','Buffer pcs','Extras pcs','Order pcs','Weight kg','Rate INR','Basis','Amount INR']]
mixed=any(line.get('uom') in ('kg','ltr') for line in po['lines'])
for line in po['lines']:rows.append([line.get('code',''),line['name'],line['mrp'],line['buffer'],line['extras'],line['orderQty'],line.get('weight'),None if po['quantityOnly'] else line.get('rate'),line.get('rateUnit',''),None if po['quantityOnly'] else line.get('amount')])
if mixed:
 rows[10]=[h.replace(' pcs',' qty') for h in rows[10]];rows[10].insert(2,'Unit')
 for row,line in zip(rows[11:],po['lines']):row.insert(2,line.get('uom','pcs'))
rows += [[],['Total before GST',None if po['quantityOnly'] else po['total']],['Notes',po.get('notes','')],['Stock allocation',po.get('stockAllocation','')]]
for row in rows:ws.append(row)
for row in ws:
 for c in row:
  if isinstance(c.value,str):c.data_type='s'
  c.alignment=Alignment(vertical='top',wrap_text=True)
  if isinstance(c.value,(float,int)):c.number_format='#,##0.00'
for rownum in [1,11]:
 for c in ws[rownum]:c.fill=PatternFill('solid',fgColor='204321');c.font=Font(color='FFFFFF',bold=True)
for i,width in enumerate([26,55]+([10] if mixed else [])+[14,14,14,14,16,16,14,20],1):ws.column_dimensions[get_column_letter(i)].width=width
if mixed:
 for row in ws.iter_rows(min_row=12,max_row=11+len(po['lines']),min_col=4,max_col=7):
  for cell in row:cell.number_format='#,##0.###'
ws.freeze_panes='D12' if mixed else 'C12';ws.auto_filter.ref=f'A11:{"K" if mixed else "J"}{11+len(po["lines"])}'
# Append identification columns so existing quantity/rate columns keep their positions.
code_col=12 if mixed else 11;photo_col=code_col+1
ws.cell(11,code_col,'Supplier part code');ws.cell(11,photo_col,'Product image')
ws.column_dimensions[get_column_letter(code_col)].width=28;ws.column_dimensions[get_column_letter(photo_col)].width=15
for rownum,line in enumerate(po['lines'],12):
 c=ws.cell(rownum,code_col,line.get('supplierPartCode') or 'Pending');c.data_type='s';c.alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[rownum].height=64
 if line.get('imageData'):
  with PILImage.open(io.BytesIO(base64.b64decode(line['imageData'],validate=True))) as decoded:
   if decoded.width*decoded.height>40000000:raise ValueError('Image is too large.')
   data=io.BytesIO();decoded.convert('RGB').save(data,format='PNG')
  image=Image(io.BytesIO(data.getvalue()));scale=min(80/image.width,72/image.height,1);image.width*=scale;image.height*=scale;ws.add_image(image,f'{get_column_letter(photo_col)}{rownum}')
 else:ws.cell(rownum,photo_col,'Photo pending')
for cell in ws[11]:cell.fill=PatternFill('solid',fgColor='204321');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
ws.auto_filter.ref=f'A11:{get_column_letter(photo_col)}{11+len(po["lines"])}';ws.print_title_rows='11:11';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
fab=wb.create_sheet('Fabrication');fab.append(['Subassembly','Drawing / code','Component','PPM pcs','kg per machine','Ordered pcs','Order kg'])
for line in po['lines']:
 if line.get('fabricated'):
  for child in line.get('children',[]):fab.append([line['name'],child.get('code') or child.get('drawingCode'),child['name'],child.get('ppm'),child.get('weight'),line['orderQty'],None if child.get('weight') is None else child['weight']*line['orderQty']])
for row in fab:
 for c in row:
  if isinstance(c.value,str):c.data_type='s'
for i,width in enumerate([32,24,48,16,20,18,20],1):fab.column_dimensions[get_column_letter(i)].width=width
fab.cell(1,8,'Supplier part code');fab.cell(1,9,'Product image');fab.column_dimensions['H'].width=28;fab.column_dimensions['I'].width=15
rownum=2
for line in po['lines']:
 if line.get('fabricated'):
  for child in line.get('children',[]):
   cell=fab.cell(rownum,8,child.get('supplierPartCode') or 'Pending');cell.data_type='s';cell.alignment=Alignment(wrap_text=True,vertical='center');fab.row_dimensions[rownum].height=64
   if child.get('imageData'):
    with PILImage.open(io.BytesIO(base64.b64decode(child['imageData'],validate=True))) as decoded:
     if decoded.width*decoded.height>40000000:raise ValueError('Image is too large.')
     data=io.BytesIO();decoded.convert('RGB').save(data,format='PNG')
    image=Image(io.BytesIO(data.getvalue()));scale=min(80/image.width,72/image.height,1);image.width*=scale;image.height*=scale;fab.add_image(image,f'I{rownum}')
   else:fab.cell(rownum,9,'Photo pending')
   rownum+=1
for row in fab:
 for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
for cell in fab[1]:cell.fill=PatternFill('solid',fgColor='204321');cell.font=Font(color='FFFFFF',bold=True)
fab.freeze_panes='D2';fab.auto_filter.ref=f'A1:I{max(1,fab.max_row)}';fab.print_title_rows='1:1';fab.sheet_properties.pageSetUpPr.fitToPage=True;fab.page_setup.orientation='landscape';fab.page_setup.fitToWidth=1;fab.page_setup.fitToHeight=0
priced_transport=[line for line in po['lines'] if line.get('transportInCost') and not po['quantityOnly']]
if priced_transport:
 basis=wb.create_sheet('Price basis');basis.append(['Component','Base INR / unit','Transport %','Transport INR / unit','Rate incl. transport INR / unit','Order quantity','Amount INR'])
 for line in priced_transport:basis.append([line['name']+' ['+line.get('uom','pcs')+']',line.get('purchaseBaseRate'),line.get('purchaseTransportPercent'),line.get('purchaseTransportAmount'),line.get('rate'),line['orderQty'],line.get('amount')])
 for row in basis:
  for c in row:
   if isinstance(c.value,str):c.data_type='s'
   c.alignment=Alignment(vertical='top',wrap_text=True)
 for i,width in enumerate([55,22,18,24,32,18,22],1):basis.column_dimensions[get_column_letter(i)].width=width
 basis.freeze_panes='B2';basis.auto_filter.ref=f'A1:G{basis.max_row}';basis.row_dimensions[1].height=36
 for c in basis[1]:c.fill=PatternFill('solid',fgColor='204321');c.font=Font(color='FFFFFF',bold=True)
 for row in basis.iter_rows(min_row=2,min_col=2):
  for c in row:c.number_format='#,##0.###' if c.column==6 else '#,##0.00'
out=io.BytesIO();wb.save(out);sys.stdout.buffer.write(out.getvalue())
