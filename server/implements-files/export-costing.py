"""Cost register with filterable model and segment detail. Values are saved calculation snapshots."""
import io,json,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter

def sheet(wb,title,headers,rows,widths,money_columns=()):
    ws=wb.create_sheet(title);ws.append(headers)
    for values in rows:
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if cell.column in money_columns and isinstance(cell.value,(int,float)):cell.number_format='#,##0.00'
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='204321');cell.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=32
    for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='C2';ws.sheet_view.showGridLines=False
    if rows:
        table=Table(displayName=title.replace(' ','')+'Register',ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}');table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);ws.add_table(table)
    ws.print_title_rows='1:1';ws.page_setup.orientation='landscape';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    return ws

def export(data):
    reports=data['reports'];selection=data.get('selection');wb=Workbook();wb.remove(wb.active)
    sheet(wb,'Model costs',['Model','Series','Size','Sales confirmed','BOM parts INR','Fabrication INR','Other costs INR','Known subtotal INR','Full cost per machine INR','Status','Pending inputs','BOM revision'],[[r['id'],r['series'],r['size'],'Yes' if r['salesConfirmed'] else 'No',r['partsTotal'],r['fabricationTotal'],r['otherTotal'],r['knownTotal'],r['total'],r['status'],r['pending'],r['revision']] for r in reports],[18,22,15,20,22,22,22,24,27,24,18,16],range(5,10))
    detail=[];segments=[]
    for r in reports:
        rows=[next(line for line in r['rows'] if line['id']==item_id) for item_id in selection['ids']] if selection else r['rows']
        for line in rows:detail.append([r['id'],line['segment'],line['code'],line['name'],line['type'],line['ppm'],line['weight'],line['rate'],line['unit'],line['amount'],line['status'],line.get('fabricationBaseRate'),line.get('fabricationTransportRate')])
        for name in dict.fromkeys(line['segment'] for line in rows):
            group=[line for line in rows if line['segment']==name];pending=sum(line['amount'] is None for line in group);known=round(sum(line['amount'] or 0 for line in group),2)
            segments.append([r['id'],name,len(group),known,None if pending else known,pending,'Filtered rows' if selection else 'Whole model'])
    sheet(wb,'Segment totals',['Model','Segment','Rows','Known subtotal INR','Complete subtotal INR','Pending rows','Scope'],segments,[18,32,12,25,25,18,22],(4,5))
    sheet(wb,'Cost detail',['Model','Segment','Tally code','Component / charge','Type','PPM / qty','kg per machine','Rate INR','Rate unit','Cost per machine INR','Status','Fabrication base INR/kg','Fabrication transport INR/kg'],detail,[18,30,22,65,26,15,20,20,16,26,40,25,28],(8,10,12,13))
    notes=wb.create_sheet('Read me');notes.column_dimensions['A'].width=115
    for text in ['Cost per machine: BOM parts plus fabricated subassembly plus other charges. INR before GST.','Model costs shows the whole model. Segment totals and Cost detail follow the exported row filters when used.','Blank numeric values are pending, not zero. Known subtotals exclude pending lines. Full cost stays blank until complete.','BOM amount = PPM x price per pc. Fabrication = combined kg per machine x (base price per kg + transport per kg), once.','A set-of-two combined weight is not multiplied by PPM again. Source discrepancies remain in the flags register.','PTO shaft, sticker kit and name plate are physical BOM items. Coating, assembly, rack stand and buffer cost are charges.','Stock, purchase buffers and extras do not alter this per-machine cost. No profit margin or tax is added.','These are calculated values at download time. Use the website to change inputs and recalculate; workbook edits do not sync.']:notes.append([text])
    out=io.BytesIO();wb.save(out);return out.getvalue()

try:sys.stdout.buffer.write(export(json.load(sys.stdin)))
except Exception as e:print(str(e),file=sys.stderr);sys.exit(1)
