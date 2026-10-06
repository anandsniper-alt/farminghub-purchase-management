"""Portable, editable item register. Images are embedded; all supplied text stays literal."""
import sys,json,io,base64
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.drawing.image import Image
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

def export(data):
    wb=Workbook();ws=wb.active;ws.title='Item Master'
    columns=[('image','Product image'),('code','Tally code (IMP-XXX)'),('name','Item name'),('category','Category'),('type','Type'),('unit','Unit'),('models','Models'),('drawing','Drawing code'),('supplier','Supplier'),('imageStatus','Image status'),('supplierPartCodes','Supplier part codes')]
    ws.append([label for _,label in columns])
    for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='204321');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=32
    for index,row in enumerate(data['rows'],2):
        for col,(key,_) in enumerate(columns,1):
            cell=ws.cell(index,col,'' if key=='image' else row.get(key,''));cell.data_type='s';cell.alignment=Alignment(wrap_text=True,vertical='center');cell.font=Font(name='Calibri',size=11,color='203723')
        ws.row_dimensions[index].height=66
        if row.get('image'):
            raw=base64.b64decode(row['image'],validate=True)
            with PILImage.open(io.BytesIO(raw)) as check:
                if check.format not in ('JPEG','PNG','WEBP') or check.width*check.height>40000000:raise ValueError('Unsupported item image.')
                check.verify()
            # Excel embeds PNG/JPEG reliably; convert the browser's compact WebP encoding.
            with PILImage.open(io.BytesIO(raw)) as decoded:
                if decoded.format=='WEBP':
                    converted=io.BytesIO();decoded.save(converted,format='PNG');raw=converted.getvalue()
            image=Image(io.BytesIO(raw));scale=min(88/image.width,76/image.height,1);image.width*=scale;image.height*=scale;ws.add_image(image,f'A{index}')
    for i,width in enumerate([15,23,62,29,22,10,45,20,30,22,45],1):ws.column_dimensions[get_column_letter(i)].width=width
    ws.freeze_panes='C2';ws.sheet_view.showGridLines=False
    table=Table(displayName='ItemMaster',ref=f'A1:K{len(data["rows"])+1}');table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True);ws.add_table(table)
    ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    notes=wb.create_sheet('Read me');notes.append(['Rotavator Item Master']);notes.append(['Downloaded from Farming Hub Implements purchase.']);notes.append(['Excel edits do not sync back. Use Item Master > Edit item to update the site.']);notes.append(['Photos are references. Confirm item codes, dimensions and grades for purchasing.']);notes.append(['Fabrication children are a component schedule; order the model subassembly through MRP.']);notes.append(['Missing Tally codes stay blank. PPM and weights remain in model BOMs.']);notes.column_dimensions['A'].width=110
    out=io.BytesIO();wb.save(out);return out.getvalue()

try:sys.stdout.buffer.write(export(json.load(sys.stdin)))
except Exception as exc:print(str(exc),file=sys.stderr);sys.exit(1)
