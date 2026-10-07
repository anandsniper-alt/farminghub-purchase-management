import sys,json,base64,io,csv,zipfile
from pathlib import Path
from xml.sax.saxutils import escape
ROOT=Path(__file__).resolve().parent

def stock(data):
    name=str(data.get('name','')).lower()
    raw=base64.b64decode(data.get('data',''),validate=True)
    if len(raw)>10*1024*1024: raise ValueError('Maximum file size is 10 MB.')
    if name.endswith('.csv'):
        text=raw.decode('utf-8-sig')
        if not text.strip(): raise ValueError('The file is empty. Include column headers and data rows.')
        delimiter='\t' if '\t' in text.splitlines()[0] else ','
        rows=list(csv.reader(io.StringIO(text),delimiter=delimiter))
        if len(rows)>20001 or any(len(r)>200 for r in rows): raise ValueError('Limit: 20,000 stock rows and 200 columns.')
        return {'sheets':[{'name':'Sheet 1','rows':rows}]}
    if not name.endswith('.xlsx'): raise ValueError('Choose an .xlsx or UTF-8 .csv stock sheet.')
    import openpyxl
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist())>100*1024*1024: raise ValueError('The expanded workbook is too large.')
    wb=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    if len(wb.worksheets)>30: raise ValueError('Limit: 30 sheets.')
    sheets=[]
    for s in wb:
        if (s.max_row or 0)>20001 or (s.max_column or 0)>200: raise ValueError(f'{s.title}: limit is 20,000 rows and 200 columns.')
        rows=[]
        for index,row in enumerate(s.iter_rows(values_only=True)):
            if index>=20001 or len(row)>200: raise ValueError(f'{s.title}: limit is 20,000 rows and 200 columns.')
            rows.append([v if isinstance(v,(str,int,float,bool)) or v is None else str(v) for v in row])
        sheets.append({'name':s.title,'rows':rows})
    wb.close()
    return {'sheets':sheets}

def pdf(po):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib.enums import TA_RIGHT
    out=io.BytesIO();styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='SmallCell',fontName='Helvetica',fontSize=8,leading=11,textColor=colors.HexColor('#203a25')))
    styles.add(ParagraphStyle(name='WhiteCell',parent=styles['SmallCell'],textColor=colors.white,fontName='Helvetica-Bold'))
    styles.add(ParagraphStyle(name='SmallNote',parent=styles['SmallCell'],fontSize=9,leading=13))
    def p(v,style='SmallCell'):return Paragraph(escape(str(v or '')).replace('\n','<br/>'),styles[style])
    def n(v):return '' if v is None else f'{v:,.3f}'.rstrip('0').rstrip('.')
    def money(v):return 'Pending' if v is None else f'{v:,.2f}'
    def product_image(line):
        if not line.get('imageData'):return p('Photo pending')
        with PillowImage.open(io.BytesIO(base64.b64decode(line['imageData'],validate=True))) as decoded:
            if decoded.width*decoded.height>40000000:raise ValueError('Image is too large.')
            data=io.BytesIO();decoded.convert('RGB').save(data,format='PNG');width,height=decoded.size
        scale=min(52/width,52/height);return Image(io.BytesIO(data.getvalue()),width=width*scale,height=height*scale)
    doc=SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=32,rightMargin=32,topMargin=30,bottomMargin=38)
    logo=ROOT/'../../web/assets/farming-hub-logo.png'
    from PIL import Image as PillowImage
    with PillowImage.open(logo) as dimensions:iw,ih=dimensions.size
    image=Image(str(logo),width=126,height=126*ih/iw)
    heading=Table([[image,p('PURCHASE ORDER\n'+po['id'],'Heading2')]],colWidths=[500,278]);heading.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP')]))
    story=[heading,Spacer(1,12),p('PO status: '+po.get('lifecycleStatus','DRAFT').replace('_',' '),'SmallNote'),Spacer(1,12)]
    supplier=po['supplier'];buyer=po['buyer']
    info=Table([[p('SUPPLIER\n'+supplier['name']+'\n'+supplier.get('address','')+'\nGSTIN: '+(supplier.get('gstin') or 'Pending')+'\n'+supplier.get('phone','')),p('BUYER\n'+buyer.get('buyer','')+'\n'+buyer.get('address','')+'\nGSTIN: '+(buyer.get('gstin') or 'Pending')),p('PO date: '+po['date']+'\nDelivery: '+(po.get('currentDelivery') or po.get('delivery') or 'To be agreed')+'\nPlan month: '+(po.get('planningMonth') or 'Not assigned')+'\nStock date: '+(po.get('stockAsOf') or 'Not supplied'))]],colWidths=[300,290,188])
    info.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#f2f5dd')),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]));story.extend([info,Spacer(1,16)])
    priced=not po['quantityOnly']
    mixed=any(l.get('uom') in ('ltr','kg') for l in po['lines'])
    heads=['Item code / component','MRP\nqty' if mixed else 'MRP\npcs','Buffer\nqty' if mixed else 'Buffer\npcs','Extras\nqty' if mixed else 'Extras\npcs','Order\nqty' if mixed else 'Order\npcs','Weight\nkg']+(['Rate INR','Basis','Amount INR'] if priced else [])
    rows=[[p(h,'WhiteCell') for h in heads]]
    for l in po['lines']:
        label=(l.get('code') or 'Code pending')+'\n'+l['name']+'\nSupplier part code: '+(l.get('supplierPartCode') or 'Pending')
        if l.get('fabricated'):label+='\nFabrication subassembly (one pc per machine)'
        if mixed:label+='\nQuantity unit: '+l.get('uom','pcs')
        if l.get('transportInCost') and priced:label+='\nBase INR '+money(l.get('purchaseBaseRate'))+' + '+n(l.get('purchaseTransportPercent'))+'% transport (included in rate)'
        product=Table([[product_image(l),p(label)]],colWidths=[56,178 if priced else 300]);product.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),3)]))
        line=[product,n(l['mrp']),n(l['buffer']),n(l['extras']),n(l['orderQty']),n(l.get('weight')) if l.get('weight') is not None else ('Pending' if l.get('fabricated') else '-')]
        if priced:line += [money(l['rate']),l['rateUnit'],money(l['amount'])]
        rows.append(line)
    widths=[278,53,53,53,58,68,75,45,95] if priced else [393,70,70,70,80,95]
    # Fit the writable page width exactly for both document types.
    scale=778/sum(widths);widths=[w*scale for w in widths]
    table=Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT')
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#204321')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7f8f1')]),('FONTNAME',(0,1),(-1,-1),'Helvetica'),('FONTSIZE',(0,1),(-1,-1),8),('ALIGN',(1,1),(-1,-1),'RIGHT'),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),('LINEBELOW',(0,0),(-1,0),0.7,colors.HexColor('#204321'))]))
    story.extend([table,Spacer(1,12),p(('Total before GST: INR '+money(po['total'])) if priced else 'Quantity-only PO. Rates, taxes and commercial totals are to be agreed.','SmallNote'),Spacer(1,8),p('MRP = combined model demand minus available stock, minimum zero. Buffer is shown separately; extras are manual additions. Missing stock is treated as zero until stock is supplied.','SmallNote')])
    if po.get('stockAllocation'):story.extend([Spacer(1,8),p('Stock allocation: '+po['stockAllocation'],'SmallNote')])
    if supplier.get('terms'):story.extend([Spacer(1,8),p('Payment terms: '+supplier['terms'],'SmallNote')])
    if po.get('notes'):story.extend([Spacer(1,8),p('Notes: '+po['notes'],'SmallNote')])
    for l in po['lines']:
        if not l.get('fabricated'):continue
        story.extend([Spacer(1,18),p(l['name']+' - component schedule','Heading3')])
        children=[[p('Code / drawing','WhiteCell'),p('Fabricated component','WhiteCell'),p('PPM (pcs)','WhiteCell'),p('kg / machine','WhiteCell'),p('Order weight kg','WhiteCell')]]
        for c in l.get('children',[]):children.append([p(c.get('code') or c.get('drawingCode') or 'Pending'),[product_image(c),p(c['name']+'\nSupplier part code: '+(c.get('supplierPartCode') or 'Pending'))],n(c.get('ppm')) if c.get('ppm') is not None else 'Pending',n(c.get('weight')) if c.get('weight') is not None else 'Pending',n(c['weight']*l['orderQty']) if c.get('weight') is not None else 'Pending'])
        t=Table(children,colWidths=[90,390,70,110,118],repeatRows=1)
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#204321')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f7f8f1')]),('FONTSIZE',(0,1),(-1,-1),8),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.append(t)
    def footer(canvas,document):
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#52664f'));canvas.drawString(32,20,po['id']+' | '+po.get('lifecycleStatus','DRAFT').replace('_',' '));canvas.drawRightString(810,20,f'Page {document.page}')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()

try:
    data=json.load(sys.stdin)
    if sys.argv[1]=='stock':print(json.dumps(stock(data),ensure_ascii=False))
    else:sys.stdout.buffer.write(pdf(data))
except Exception as e:
    print(str(e),file=sys.stderr);sys.exit(1)
