"""Read-only validation of the generated XLSX, independent of the template filler."""
import io
import sys
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

data = sys.stdin.buffer.read()
mode = sys.argv[1]
book = load_workbook(io.BytesIO(data), data_only=False)
cached = load_workbook(io.BytesIO(data), data_only=True)
assert book.sheetnames == ['Summary', 'Expense Breakdown', 'Item Costs']
s, e, i = [book[n] for n in book.sheetnames]
cs, ce, ci = [cached[n] for n in book.sheetnames]
assert s['C11'].value.year == 2026
assert i.freeze_panes == 'D6'
assert not s.sheet_view.showGridLines
if mode == 'final':
    assert s['C8'].value == '00012' and s['C9'].value == '00042'
    assert s['F5'].value == 1006.01 and ce['C16'].value == 1006.01
    assert s['F9'].value == 123.45
    assert i['E5'].value == 'Invoice unit price (RMB)'
    assert ci['F8'].value == 500.02 and ci['I8'].value == 1006.01
    assert ci['D8'].value == 5 and i['H6'].number_format == '#,##0.0000'
    assert i['I8'].data_type == 'f' and e['C16'].data_type == 'f'
    assert e['C26'].value == 150 and e['D26'].value == 123.45
    assert s['F18'].value == 7
elif mode == 'pending':
    assert s['C12'].value == 'PROVISIONAL'
    assert s['F5'].value is None and s['F9'].value is None
    assert e['C10'].value is None and e['D10'].value == 'Pending'
    assert e['C9'].value == 0 and e['D9'].value == 'Recorded'
    assert ce['C16'].value is None and ci['I8'].value is None
elif mode == 'literal':
    assert s['C5'].value == '=1+1' and s['C5'].data_type == 's'
    assert s['C9'].value == '@SUM(A1:A2)' and s['C9'].data_type == 's'
    assert s['C8'].value == '+0012'
    assert i['C6'].value.startswith('=HYPERLINK') and i['C6'].data_type == 's'
elif mode == 'large':
    assert i['B505'].value == 'SKU-499' and i['B506'].value == 'Total'
    assert ci['D506'].value == 1000
    assert i.auto_filter.ref == 'B5:I505'
    assert 'I505' in i['I506'].value
with zipfile.ZipFile(io.BytesIO(data)) as z:
    for name in z.namelist():
        if name.endswith('.xml'):
            ET.fromstring(z.read(name))
print(mode + ' verified')
